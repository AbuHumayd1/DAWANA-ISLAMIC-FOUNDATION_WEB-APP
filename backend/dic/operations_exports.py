import io
from pathlib import Path

from django.utils import timezone

from .models import HouseAllocation, PaymentAttempt, RegistrationOrder
from .pricing import (
    PricingConfigurationError,
    current_order_prices,
    revision_matches_current,
)

def spreadsheet_safe(value):
    text = "" if value is None else str(value)
    if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(
        ("\t", "\r", "\n")
    ):
        return "'" + text
    return text


def registration_export_rows(queryset, include_contacts=False):
    headers = [
        "order_reference",
        "edition",
        "submitted_at",
        "order_status",
        "payment_status",
        "verified_payment",
        "total_amount_ngn",
        "participant_count",
        "categories",
        "houses",
        "repricing_status",
    ]
    if include_contacts:
        headers.extend(("participant_names", "participant_emails", "participant_phones"))

    rows = []
    if hasattr(queryset, "select_related"):
        queryset = queryset.select_related(
            "edition",
            "primary_registrant",
            "current_price_revision",
        ).prefetch_related(
            "items__participant",
            "items__participant__house_allocations__house",
            "items__participant__house_allocations__edition",
            "payment_attempts",
            "current_price_revision__items",
        )
    for order in queryset:
        participants = list(order.items.all())
        verified = (
            order.status == RegistrationOrder.Status.CONFIRMED
            and order.payment_status == RegistrationOrder.PaymentStatus.PAID
            and any(
                attempt.status == PaymentAttempt.Status.SUCCESSFUL
                for attempt in order.payment_attempts.all()
            )
        )
        houses = sorted(
            {
                allocation.house.name
                for item in participants
                for allocation in item.participant.house_allocations.all()
                if allocation.edition_id == order.edition_id
                and allocation.status == HouseAllocation.Status.ALLOCATED
            }
        )
        try:
            current_prices = current_order_prices(order)
            repricing_required = order.payment_status != RegistrationOrder.PaymentStatus.PAID and not revision_matches_current(
                order.current_price_revision,
                current_prices,
            )
        except PricingConfigurationError:
            repricing_required = order.payment_status != RegistrationOrder.PaymentStatus.PAID
        row = [
            str(order.public_id),
            order.edition.name,
            order.created_at.isoformat(),
            order.status,
            order.payment_status,
            "verified" if verified else "not verified",
            order.total_amount_ngn,
            len(participants),
            ", ".join(sorted({item.category_name_snapshot for item in participants})),
            ", ".join(houses),
            "Required" if repricing_required else "Current",
        ]
        if include_contacts:
            row.extend(
                (
                    ", ".join(item.participant.full_name for item in participants),
                    ", ".join(item.participant.email for item in participants),
                    ", ".join(item.participant.phone for item in participants),
                )
            )
        rows.append(row)
    return headers, rows


def create_word_report(title, edition_name, filters, headers, rows):
    try:
        from docx import Document
        from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.section import WD_ORIENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt, RGBColor
    except ImportError as exc:
        raise RuntimeError("python-docx is required for Word exports.") from exc

    document = Document()
    section = document.sections[0]
    if len(headers) > 7:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(45, 55, 52)
    section.header.paragraphs[0].text = "DAWANA ISLAMIC FOUNDATION | DIC Operations"
    logo_path = Path(__file__).resolve().parent / "assets" / "dawana-logo.png"
    if logo_path.is_file():
        section.header.paragraphs[0].add_run().add_picture(str(logo_path), width=Inches(0.4))
    heading = document.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run("DAWANA ISLAMIC FOUNDATION")
    run.bold = True
    run.font.name = "Aptos Display"
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor(15, 77, 56)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0, 53, 36)
    document.add_paragraph(f"Edition: {edition_name or 'All editions'}")
    document.add_paragraph(
        f"Generated: {timezone.localtime().strftime('%d %B %Y, %H:%M %Z')}"
    )
    document.add_paragraph(f"Filters: {filters or 'None'}")
    document.add_paragraph(f"Records: {len(rows)}")
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    header_cells = table.rows[0].cells
    for index, label in enumerate(headers):
        header_cells[index].text = str(label)
        header_cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in header_cells[index].paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "0F4D38")
        header_cells[index]._tc.get_or_add_tcPr().append(shading)
    repeat_header = OxmlElement("w:tblHeader")
    repeat_header.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat_header)
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = str(value)
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Dawana Islamic Foundation — Confidential operational report")
    page_field = OxmlElement("w:fldSimple")
    page_field.set(qn("w:instr"), "PAGE")
    footer._p.append(page_field)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
