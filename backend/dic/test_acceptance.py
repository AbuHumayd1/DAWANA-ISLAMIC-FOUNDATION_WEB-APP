"""Acceptance fixtures run only inside Django's transactional test database."""
import csv
import io
import json
import os
from pathlib import Path

from django.contrib.auth.models import Permission
from django.test import TestCase
from django.urls import reverse
from docx import Document
from docx.oxml.ns import qn

from .models import House, RegistrationCategory, RegistrationOrder
from . import test_operations_portal as portal_fixtures


class OperationsAcceptanceTests(TestCase):
    setUp = portal_fixtures.DICOperationsPortalTests.setUp
    create_user = portal_fixtures.DICOperationsPortalTests.create_user
    make_order = portal_fixtures.DICOperationsPortalTests.make_order

    def test_representative_reports_and_export_permissions(self):
        # Deliberately long, explicitly synthetic labels exercise wrapping.
        self.edition.name = "Acceptance Test Camp 2097 — Long Text and Multipage Report Review"
        self.edition.save()
        for index in range(32):
            self.make_order((f"Acceptance Sample {index:02d} Participant With A Deliberately Long Name",), paid=index % 2 == 0)
            RegistrationCategory.objects.create(edition=self.edition, name=f"Acceptance Sample {index:02d} Category With A Long Descriptive Label", fee_ngn=125)
            House.objects.create(edition=self.edition, name=f"Acceptance Sample {index:02d} House With A Long Descriptive Label", gender="male")
        self.finance.user_permissions.add(*Permission.objects.filter(content_type__app_label="dic", codename__in=("can_export_dic_contact_information", "can_view_dic_contact_information")))
        self.client.force_login(self.finance)
        output = os.environ.get("DIC_ACCEPTANCE_REPORT_DIR")
        output_path = Path(output).resolve() if output else None
        if output_path:
            output_path.mkdir(parents=True, exist_ok=True)
        manifest = {}
        for report in ("registrations", "participants", "payments", "categories", "houses", "attention"):
            responses = {}
            for export_format in ("csv", "docx"):
                response = self.client.get(reverse("dic:operations-report"), {"edition": self.edition.pk, "report": report, "format": export_format})
                self.assertEqual(response.status_code, 200, (report, export_format))
                responses[export_format] = response.content
                if output_path:
                    (output_path / f"acceptance-{report}.{export_format}").write_bytes(response.content)
            rows = list(csv.reader(io.StringIO(responses["csv"].decode("utf-8-sig"))))
            document = Document(io.BytesIO(responses["docx"]))
            self.assertEqual(len(document.tables[0].rows), len(rows))
            self.assertIn("DAWANA ISLAMIC FOUNDATION", document.paragraphs[0].text)
            self.assertIsNotNone(document.tables[0].rows[0]._tr.find(".//" + qn("w:tblHeader")))
            self.assertIn("PAGE", document.sections[0].footer._element.xml)
            section = document.sections[0]
            self.assertEqual(section.page_width > section.page_height, len(rows[0]) > 7)
            self.assertGreaterEqual(len(rows) - 1, 16)
            manifest[report] = {"records": len(rows) - 1, "columns": len(rows[0]), "orientation": "landscape" if section.page_width > section.page_height else "portrait", "visual_rendering": "not verified"}
        if output_path:
            (output_path / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
        # Ordinary operators cannot obtain finance reports in either format.
        self.client.force_login(self.operator)
        for export_format in ("csv", "docx"):
            self.assertEqual(self.client.get(reverse("dic:operations-report"), {"report": "payments", "format": export_format}).status_code, 403)
        permissionless = self.create_user("acceptance-no-export", "test-only", "can_view_dic_operations")
        self.client.force_login(permissionless)
        for export_format in ("csv", "docx"):
            self.assertEqual(self.client.get(reverse("dic:operations-report"), {"report": "registrations", "format": export_format}).status_code, 403)

    def test_empty_edition_returns_zero_metrics_and_empty_directories(self):
        RegistrationOrder.objects.filter(edition=self.edition).update(edition=self._empty_edition().pk)
        self.client.force_login(self.operator)
        data = self.client.get(reverse("dic:operations-dashboard"), {"edition": self.edition.pk}).json()
        self.assertEqual(data["metrics"]["orders"], 0)
        self.assertEqual(data["metrics"]["participants"], 0)
        for name in ("operations-registrations", "operations-participants"):
            response = self.client.get(reverse(f"dic:{name}"), {"edition": self.edition.pk})
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["results"], [])

    def _empty_edition(self):
        from .models import CampEdition
        return CampEdition.objects.create(name="Acceptance Other Edition", year=2098, camp_start_date=self.edition.camp_start_date, camp_end_date=self.edition.camp_end_date)

    def test_expired_session_loses_access(self):
        from django.contrib.sessions.models import Session
        from django.utils import timezone
        from datetime import timedelta
        self.client.force_login(self.operator)
        Session.objects.filter(session_key=self.client.session.session_key).update(expire_date=timezone.now() - timedelta(seconds=1))
        self.assertEqual(self.client.get(reverse("dic:operations-me")).status_code, 403)
